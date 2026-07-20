"""The phase's actual acceptance test: the same result set produces
consistent counts across all four report formats.

Each format is re-parsed with an appropriate reader (json, csv,
openpyxl) or a targeted marker count (HTML), and the extracted "number
of individual findings" and "per-severity breakdown" are compared
against each other and against a hand-computed expectation — not just
asserted to be "non-zero" or "present."
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl

from insightengine.core.types import Severity
from insightengine.reporting.csv_report import CSVReportGenerator
from insightengine.reporting.excel_report import ExcelReportGenerator
from insightengine.reporting.html_report import HTMLReportGenerator
from insightengine.reporting.json_report import JSONReportGenerator
from insightengine.reporting.models import build_report_data
from insightengine.validation.results import ValidationResult

_RESULTS = (
    ValidationResult(rule_id="age_range", severity=Severity.ERROR, message="too young", row_index=0),
    ValidationResult(rule_id="age_range", severity=Severity.ERROR, message="too old", row_index=3),
    ValidationResult(rule_id="q2_gate", severity=Severity.WARNING, message="q2 missing", row_index=1),
    ValidationResult(rule_id="q2_gate", severity=Severity.WARNING, message="q2 missing", row_index=2),
    ValidationResult(rule_id="q2_gate", severity=Severity.WARNING, message="q2 missing", row_index=5),
    ValidationResult(rule_id="sum100", severity=Severity.INFO, message="fyi", row_index=4),
)
# Hand-computed expectations: 6 total, error=2, warning=3, info=1.
_EXPECTED_TOTAL = 6
_EXPECTED_BY_SEVERITY = {"error": 2, "warning": 3, "info": 1}


def _count_json(path: Path) -> tuple[int, dict[str, int]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    assert payload["summary"]["total"] == len(payload["results"])
    return len(payload["results"]), payload["summary"]["by_severity"]


def _count_csv(path: Path) -> tuple[int, dict[str, int]]:
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    total = len(rows)
    by_severity: dict[str, int] = {}
    for row in rows:
        by_severity[row["severity"]] = by_severity.get(row["severity"], 0) + 1
    return total, by_severity


def _count_html(path: Path) -> tuple[int, dict[str, int]]:
    content = path.read_text(encoding="utf-8")
    # Result rows (one per finding) only ever appear in the Details
    # section - the Summary section's rows are one per *rule*, and reuse
    # the same severity CSS classes/text, so counting across the whole
    # document would double-count severities from both sections.
    details_section = content.split("<h2>Details</h2>", 1)[1]
    total = details_section.count('class="result-row"')
    by_severity = {
        "error": details_section.count('severity-error">error'),
        "warning": details_section.count('severity-warning">warning'),
        "info": details_section.count('severity-info">info'),
    }
    return total, by_severity


def _count_excel(path: Path) -> tuple[int, dict[str, int]]:
    workbook = openpyxl.load_workbook(path)
    details = workbook["Details"]
    total = details.max_row - 1
    by_severity: dict[str, int] = {}
    for row in details.iter_rows(min_row=2, values_only=True):
        severity = row[1]
        by_severity[severity] = by_severity.get(severity, 0) + 1
    return total, by_severity


class TestAcceptanceCriterion:
    def test_all_four_formats_report_consistent_counts(self, tmp_path: Path) -> None:
        data = build_report_data(_RESULTS)

        json_path = JSONReportGenerator().build(data, tmp_path / "report.json")
        csv_path = CSVReportGenerator().build(data, tmp_path / "report.csv")
        html_path = HTMLReportGenerator().build(data, tmp_path / "report.html")
        excel_path = ExcelReportGenerator().build(data, tmp_path / "report.xlsx")

        json_total, json_by_severity = _count_json(json_path)
        csv_total, csv_by_severity = _count_csv(csv_path)
        html_total, html_by_severity = _count_html(html_path)
        excel_total, excel_by_severity = _count_excel(excel_path)

        # Every format agrees on the hand-computed expectation...
        assert json_total == _EXPECTED_TOTAL
        assert csv_total == _EXPECTED_TOTAL
        assert html_total == _EXPECTED_TOTAL
        assert excel_total == _EXPECTED_TOTAL

        # ...not just on each other by coincidence (e.g. all four being
        # wrong in the same way would still pass a cross-format-only check).
        assert json_by_severity == _EXPECTED_BY_SEVERITY
        assert {k: v for k, v in csv_by_severity.items()} == _EXPECTED_BY_SEVERITY
        assert html_by_severity == _EXPECTED_BY_SEVERITY
        assert excel_by_severity == _EXPECTED_BY_SEVERITY

    def test_empty_result_set_consistent_across_formats(self, tmp_path: Path) -> None:
        data = build_report_data([])

        json_path = JSONReportGenerator().build(data, tmp_path / "report.json")
        csv_path = CSVReportGenerator().build(data, tmp_path / "report.csv")
        html_path = HTMLReportGenerator().build(data, tmp_path / "report.html")
        excel_path = ExcelReportGenerator().build(data, tmp_path / "report.xlsx")

        assert _count_json(json_path)[0] == 0
        assert _count_csv(csv_path)[0] == 0
        assert _count_html(html_path)[0] == 0
        assert _count_excel(excel_path)[0] == 0
