"""Unit tests for insightengine.reporting.excel_report."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from insightengine.core.types import Severity
from insightengine.reporting.excel_report import ExcelReportGenerator
from insightengine.reporting.models import build_report_data
from insightengine.validation.results import ValidationResult


class TestExcelReportGenerator:
    def test_creates_parent_directories(self, tmp_path: Path) -> None:
        data = build_report_data([])
        path = tmp_path / "nested" / "report.xlsx"
        ExcelReportGenerator().build(data, path)
        assert path.is_file()

    def test_creates_summary_and_details_sheets(self, tmp_path: Path) -> None:
        path = tmp_path / "report.xlsx"
        ExcelReportGenerator().build(build_report_data([]), path)
        workbook = openpyxl.load_workbook(path)
        assert workbook.sheetnames == ["Summary", "Details"]

    def test_details_sheet_row_count_matches_results(self, tmp_path: Path) -> None:
        results = [
            ValidationResult(rule_id="r1", severity=Severity.ERROR, message="x", row_index=i)
            for i in range(5)
        ]
        data = build_report_data(results)
        path = tmp_path / "report.xlsx"
        ExcelReportGenerator().build(data, path)
        workbook = openpyxl.load_workbook(path)
        details = workbook["Details"]
        assert details.max_row - 1 == 5  # minus header

    def test_details_header(self, tmp_path: Path) -> None:
        path = tmp_path / "report.xlsx"
        ExcelReportGenerator().build(build_report_data([]), path)
        workbook = openpyxl.load_workbook(path)
        details = workbook["Details"]
        header = [cell.value for cell in next(details.iter_rows(min_row=1, max_row=1))]
        assert header == ["Rule", "Severity", "Message", "Row", "Variables"]

    def test_summary_sheet_contains_total(self, tmp_path: Path) -> None:
        results = [
            ValidationResult(rule_id="r1", severity=Severity.ERROR, message="x", row_index=i)
            for i in range(3)
        ]
        data = build_report_data(results)
        path = tmp_path / "report.xlsx"
        ExcelReportGenerator().build(data, path)
        workbook = openpyxl.load_workbook(path)
        summary = workbook["Summary"]
        assert summary["A2"].value == "Total findings"
        assert summary["B2"].value == 3

    def test_result_field_values_present_in_details(self, tmp_path: Path) -> None:
        results = [
            ValidationResult(
                rule_id="r1",
                severity=Severity.ERROR,
                message="bad value",
                row_index=2,
                variables=("Q1", "Q2"),
            )
        ]
        data = build_report_data(results)
        path = tmp_path / "report.xlsx"
        ExcelReportGenerator().build(data, path)
        workbook = openpyxl.load_workbook(path)
        details = workbook["Details"]
        row = [cell.value for cell in next(details.iter_rows(min_row=2, max_row=2))]
        assert row == ["r1", "error", "bad value", 2, "Q1, Q2"]
